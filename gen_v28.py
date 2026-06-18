#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""粤闽浙核心业绩看板 v28 - Excel表格版"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import date
import calendar

TODAY = date(2026, 5, 30)
_last_day = calendar.monthrange(TODAY.year, TODAY.month)[1]
DAYS_LEFT = (_last_day - TODAY.day) + 1

df = pd.read_excel('/tmp/files/1780106178.xlsx', header=None)

op_map = {
    '泉州市小花猫-泉州乐生活组': '泉州小花猫',
    '慈溪哲光-金华义乌乐生活组': '慈溪哲光',
    '惠州市巨浪-惠州乐生活组': '惠州巨浪',
    '义乌市汐洋-金华义乌乐生活组': '义乌汐洋',
    '惠安广通-晋江乐生活组': '惠安广通',
    '惠州市柏万-惠阳乐生活组': '惠州柏万',
}

OPS = []
for i, row in df.iterrows():
    level = str(row[0]) if pd.notna(row[0]) else ""
    if level.startswith('L2'):
        raw = str(row[1]).replace('L2:','').strip()
        name = op_map.get(raw, raw)
        def v(col):
            r = row[col]
            return r if pd.notna(r) else 0
        OPS.append({
            'name': name,
            'owner': str(row[3]) if pd.notna(row[3]) else '',
            'gtv': v(22),
            'gtv_tgt': v(24),
            'gtv_pct': v(25),
            'gtv_gap': v(26),
            'excl': v(36),
            'cpc': v(29),
            'cpc_tgt': v(31),
            'cpc_pct': v(32),
            'cpc_gap': v(33),
            'cpc_newsgn': v(38),
            'cpc_new_tgt': v(40),
            'wpu': v(34),
            'mem': v(43),
            'mem_renewal': v(14),
            'mem_renewal_mom': v(15),
            'dongxiao': v(27),
            'dongxiao_mom': v(28),
        })

total_gtv = sum(o['gtv'] for o in OPS)
total_gtv_tgt = sum(o['gtv_tgt'] for o in OPS)
total_cpc = sum(o['cpc'] for o in OPS)
total_cpc_tgt = sum(o['cpc_tgt'] for o in OPS)
total_new = sum(o['cpc_newsgn'] for o in OPS)
total_wpu = sum(o['wpu'] for o in OPS)
total_mem = sum(o['mem'] for o in OPS)
total_dongxiao = sum(o['dongxiao'] for o in OPS)

# ── 样式 ──
DARK_BLUE = PatternFill('solid', fgColor='1F4E79')
DARK_BLUE_FONT = Font(name='微软雅黑', bold=True, color='FFFFFF', size=12)
HEADER_FILL = PatternFill('solid', fgColor='D6E4F0')
HEADER_FONT = Font(name='微软雅黑', bold=True, color='1F4E79', size=10)
WHITE_FONT = Font(name='微软雅黑', size=10)
FOOTER_FILL = PatternFill('solid', fgColor='1F4E79')
FOOTER_FONT = Font(name='微软雅黑', bold=True, color='FFFFFF', size=10)
THIN_BORDER = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='thin', color='B0C4DE'),
)
CENTER = Alignment(horizontal='center', vertical='center')

def bar_color(p):
    if p >= 0.90: return '22C55E'
    if p >= 0.80: return 'FACC15'
    if p >= 0.70: return 'FB923C'
    return 'F87171'

def bar_fill(p):
    return PatternFill('solid', fgColor=bar_color(p))

def bar_font(p):
    return Font(name='微软雅黑', size=10, bold=True, color='1F4E79' if bar_color(p) in ('22C55E','FACC15') else 'FFFFFF')

def pct_color(p):
    if p >= 0.90: return '22C55E'
    if p >= 0.80: return 'FACC15'
    if p >= 0.70: return 'FB923C'
    return 'F87171'

def pct_font(p):
    return Font(name='微软雅黑', size=10, bold=True, color=pct_color(p))

wb = Workbook()

# ═══════════════ Sheet1: GTV ═══════════════
ws1 = wb.active
ws1.title = "GTV月度进度"

ws1.merge_cells('A1:I1')
ws1['A1'] = f"📊 粤闽浙城市群 GTV 月度进度（{TODAY.strftime('%Y年%m月%d日')}，剩余{DAYS_LEFT}天）"
ws1['A1'].font = Font(name='微软雅黑', bold=True, size=14, color='1F4E79')
ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 30

gtv_hdrs = ["代运营商","负责人","本月GTV(元)","月度目标(元)","GTV进度","进度%","缺口值(元)","日缺口(元)","剔除率"]
for ci, h in enumerate(gtv_hdrs, 1):
    cell = ws1.cell(row=2, column=ci, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER

for ri, o in enumerate(OPS, 3):
    p = o['gtv_pct']
    daily_gap = o['gtv_gap'] / DAYS_LEFT if DAYS_LEFT > 0 else 0
    vals = [o['name'], o['owner'], o['gtv'], o['gtv_tgt'], p, p, o['gtv_gap'], daily_gap, o['excl']]
    for ci, v in enumerate(vals, 1):
        cell = ws1.cell(row=ri, column=ci, value=v)
        cell.font = WHITE_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        if ri % 2 == 0:
            cell.fill = PatternFill('solid', fgColor='F8F9FA')
        if ci in (3, 4, 7, 8):
            cell.number_format = '#,##0'
        if ci == 5:
            cell.value = f"{p*100:.1f}%"
            cell.fill = bar_fill(p)
            cell.font = bar_font(p)
        if ci == 6:
            cell.font = pct_font(p)
        if ci == 9:
            cell.number_format = '0.0%'

ri = len(OPS) + 3
p_tot = total_gtv / total_gtv_tgt if total_gtv_tgt else 0
total_gap = total_gtv_tgt - total_gtv
daily_gap_tot = total_gap / DAYS_LEFT if DAYS_LEFT > 0 else 0
vals = ["合计", "—", total_gtv, total_gtv_tgt, p_tot, p_tot, total_gap, daily_gap_tot, "—"]
for ci, v in enumerate(vals, 1):
    cell = ws1.cell(row=ri, column=ci, value=v)
    cell.fill = FOOTER_FILL
    cell.font = FOOTER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER
    if ci in (3, 4, 7, 8):
        cell.number_format = '#,##0'
    if ci == 5:
        cell.value = f"{p_tot*100:.1f}%"
        cell.fill = bar_fill(p_tot)
        cell.font = bar_font(p_tot)
    if ci == 6:
        cell.font = pct_font(p_tot)

ws1.column_dimensions['A'].width = 14
ws1.column_dimensions['B'].width = 10
ws1.column_dimensions['C'].width = 18
ws1.column_dimensions['D'].width = 18
ws1.column_dimensions['E'].width = 12
ws1.column_dimensions['F'].width = 10
ws1.column_dimensions['G'].width = 16
ws1.column_dimensions['H'].width = 14
ws1.column_dimensions['I'].width = 10

# ═══════════════ Sheet2: CPC ═══════════════
ws2 = wb.create_sheet("CPC月度进度")
ws2.merge_cells('A1:I1')
ws2['A1'] = f"📱 粤闽浙城市群 CPC 月度进度（{TODAY.strftime('%Y年%m月%d日')}，剩余{DAYS_LEFT}天）"
ws2['A1'].font = Font(name='微软雅黑', bold=True, size=14, color='1F4E79')
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 30

cpc_hdrs = ["代运营商","负责人","CPC消耗(不含区县)","月度目标","CPC进度","进度%","缺口值","有效新签","旺铺通(元)"]
for ci, h in enumerate(cpc_hdrs, 1):
    cell = ws2.cell(row=2, column=ci, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER

for ri, o in enumerate(OPS, 3):
    p = o['cpc_pct']
    new_str = f"{o['cpc_newsgn']}/{o['cpc_new_tgt']}" if o['cpc_new_tgt'] else str(o['cpc_newsgn'])
    vals = [o['name'], o['owner'], o['cpc'], o['cpc_tgt'], p, p, o['cpc_gap'], o['cpc_newsgn'], o['wpu']]
    for ci, v in enumerate(vals, 1):
        cell = ws2.cell(row=ri, column=ci, value=v)
        cell.font = WHITE_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        if ri % 2 == 0:
            cell.fill = PatternFill('solid', fgColor='F8F9FA')
        if ci in (3, 4, 7, 9):
            cell.number_format = '#,##0'
        if ci == 5:
            cell.value = f"{p*100:.1f}%"
            cell.fill = bar_fill(p)
            cell.font = bar_font(p)
        if ci == 6:
            cell.font = pct_font(p)

ri = len(OPS) + 3
p_tot = total_cpc / total_cpc_tgt if total_cpc_tgt else 0
total_cpc_gap = total_cpc_tgt - total_cpc
vals = ["合计", "—", total_cpc, total_cpc_tgt, p_tot, p_tot, total_cpc_gap, total_new, total_wpu]
for ci, v in enumerate(vals, 1):
    cell = ws2.cell(row=ri, column=ci, value=v)
    cell.fill = FOOTER_FILL
    cell.font = FOOTER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER
    if ci in (3, 4, 7, 9):
        cell.number_format = '#,##0'
    if ci == 5:
        cell.value = f"{p_tot*100:.1f}%"
        cell.fill = bar_fill(p_tot)
        cell.font = bar_font(p_tot)
    if ci == 6:
        cell.font = pct_font(p_tot)

ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 10
ws2.column_dimensions['C'].width = 18
ws2.column_dimensions['D'].width = 16
ws2.column_dimensions['E'].width = 12
ws2.column_dimensions['F'].width = 10
ws2.column_dimensions['G'].width = 14
ws2.column_dimensions['H'].width = 12
ws2.column_dimensions['I'].width = 14

# ═══════════════ Sheet3: MEM ═══════════════
ws3 = wb.create_sheet("MEM收款与续约率")
ws3.merge_cells('A1:F1')
ws3['A1'] = f"💰 粤闽浙城市群 MEM 收款与续约率（{TODAY.strftime('%Y年%m月%d日')}）"
ws3['A1'].font = Font(name='微软雅黑', bold=True, size=14, color='1F4E79')
ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 30

mem_hdrs = ["代运营商","负责人","MEM收款额(元)","MEM续约率","续约率日变化","优质动销门店数"]
for ci, h in enumerate(mem_hdrs, 1):
    cell = ws3.cell(row=2, column=ci, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER

for ri, o in enumerate(OPS, 3):
    mom_str = f"▲{o['mem_renewal_mom']*100:.1f}%" if o['mem_renewal_mom'] >= 0 else f"▼{abs(o['mem_renewal_mom'])*100:.1f}%"
    vals = [o['name'], o['owner'], o['mem'], o['mem_renewal'], mom_str, o['dongxiao']]
    for ci, v in enumerate(vals, 1):
        cell = ws3.cell(row=ri, column=ci, value=v)
        cell.font = WHITE_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        if ri % 2 == 0:
            cell.fill = PatternFill('solid', fgColor='F8F9FA')
        if ci == 3:
            cell.number_format = '#,##0'
        if ci == 4:
            cell.number_format = '0.0%'
            if o['mem_renewal'] >= 0.65:
                cell.font = Font(name='微软雅黑', size=10, bold=True, color='22C55E')
            elif o['mem_renewal'] >= 0.55:
                cell.font = Font(name='微软雅黑', size=10, bold=True, color='FACC15')
            else:
                cell.font = Font(name='微软雅黑', size=10, bold=True, color='F87171')
        if ci == 6:
            cell.number_format = '#,##0'
            if o['dongxiao'] >= 150:
                cell.font = Font(name='微软雅黑', size=10, bold=True, color='22C55E')
            elif o['dongxiao'] >= 120:
                cell.font = Font(name='微软雅黑', size=10, bold=True, color='FACC15')
            else:
                cell.font = Font(name='微软雅黑', size=10, bold=True, color='F87171')

ri = len(OPS) + 3
vals = ["合计", "—", total_mem, "—", "—", total_dongxiao]
for ci, v in enumerate(vals, 1):
    cell = ws3.cell(row=ri, column=ci, value=v)
    cell.fill = FOOTER_FILL
    cell.font = FOOTER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER
    if ci in (3, 6):
        cell.number_format = '#,##0'

ws3.column_dimensions['A'].width = 14
ws3.column_dimensions['B'].width = 10
ws3.column_dimensions['C'].width = 18
ws3.column_dimensions['D'].width = 14
ws3.column_dimensions['E'].width = 16
ws3.column_dimensions['F'].width = 16

out = "/tmp/粤闽浙核心业绩看板_2026-05-30.xlsx"
wb.save(out)
print(f"OK: {out}")

for o in OPS:
    print(f"  {o['name']}({o['owner']}): GTV={o['gtv']:,.0f}({o['gtv_pct']*100:.1f}%) CPC={o['cpc']:,.0f}({o['cpc_pct']*100:.1f}%) MEM={o['mem']:,.0f} 续约率={o['mem_renewal']*100:.1f}% 动销={o['dongxiao']}")
print(f"  合计: GTV={total_gtv:,.0f}({total_gtv/total_gtv_tgt*100:.1f}%) CPC={total_cpc:,.0f}({total_cpc/total_cpc_tgt*100:.1f}%) MEM={total_mem:,.0f} 动销={total_dongxiao}")
