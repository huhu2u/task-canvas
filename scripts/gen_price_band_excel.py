import pandas as pd
import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys, glob, os
# 支持命令行传参：python3 gen_price_band_excel.py <文件路径>
if len(sys.argv) > 1:
    _input_file = sys.argv[1]
else:
    # 自动找最新的超价格带文件
    _candidates = sorted(glob.glob('/tmp/files/price_band_*.xlsx') + glob.glob('/tmp/files/*超价格带*.xlsx'), key=os.path.getmtime)
    _input_file = _candidates[-1] if _candidates else '/tmp/files/price_band_latest.xlsx'
df = pd.read_excel(_input_file, sheet_name=0)

def get_distributor(org):
    org = str(org)
    if '巨浪' in org: return '惠州巨浪'
    if '柏万' in org: return '惠州柏万'
    if '小花猫' in org: return '泉州小花猫'
    if '广通' in org: return '惠安广通'
    if '哲光' in org: return '慈溪哲光'
    if '汐洋' in org: return '义乌汐洋'
    return '其他'

df['代运营商'] = df['销售业绩归属组织名称'].apply(get_distributor)
df_6 = df[df['代运营商'] != '其他'].copy()

BAND_MAP = {300: ('品类A', '≤300元', 0.17), 500: ('品类B', '≤500元', 0.34), 1000: ('品类C', '≤1000元', 0.06)}
DIST_ORDER = ['惠州巨浪','惠州柏万','泉州小花猫','惠安广通','慈溪哲光','义乌汐洋']

rows = []
for band, (bname, brange, threshold) in BAND_MAP.items():
    sub = df_6[df_6['门店所属订单价格带'] == band]
    gtv = sub['实付验证GTV'].sum()
    over_gtv = sub['超价格带订单产生的实付验证GTV'].sum()
    shops = sub['门店ID'].nunique()
    ratio = over_gtv / gtv if gtv > 0 else 0
    status = '✅ 达标' if ratio <= threshold else '❌ 超标'
    diff = ratio - threshold
    rows.append(['粤闽浙合计', f'{bname}({brange})', f'{threshold*100:.0f}%',
                 round(gtv,2), round(over_gtv,2), shops,
                 f'{ratio*100:.1f}%', status, f'{diff*100:+.1f}%', bname, True])
    for dist in DIST_ORDER:
        sub2 = sub[sub['代运营商'] == dist]
        gtv2 = sub2['实付验证GTV'].sum()
        over2 = sub2['超价格带订单产生的实付验证GTV'].sum()
        shops2 = sub2['门店ID'].nunique()
        ratio2 = over2 / gtv2 if gtv2 > 0 else 0
        status2 = '✅ 达标' if ratio2 <= threshold else '❌ 超标'
        diff2 = ratio2 - threshold
        rows.append([dist, f'{bname}({brange})', f'{threshold*100:.0f}%',
                     round(gtv2,2), round(over2,2), shops2,
                     f'{ratio2*100:.1f}%', status2, f'{diff2*100:+.1f}%', bname, False])

df_s1 = pd.DataFrame(rows, columns=['代运营商','价格带','阈值','实付GTV','超价格带GTV','门店数','实际占比','达标状态','超出/剩余','_band','_is_total'])

BAND_COLORS = {
    '品类A': {'total_bg': 'FF1A5276', 'total_fg': 'FFFFFFFF', 'row_bg': 'FFD6EAF8', 'row_fg': 'FF1A5276'},
    '品类B': {'total_bg': 'FF784212', 'total_fg': 'FFFFFFFF', 'row_bg': 'FFFDEBD0', 'row_fg': 'FF784212'},
    '品类C': {'total_bg': 'FF4A235A', 'total_fg': 'FFFFFFFF', 'row_bg': 'FFF5EEF8', 'row_fg': 'FF4A235A'},
}

CLR_HEADER  = 'FF2C3E50'
CLR_RED_BG  = 'FFFEE2E2'; CLR_RED_FG  = 'FFB91C1C'
CLR_YLW_BG  = 'FFFEF9C3'; CLR_YLW_FG  = 'FF92400E'
CLR_GRN_BG  = 'FFDCFCE7'; CLR_GRN_FG  = 'FF166534'

def make_fill(hex6): return PatternFill('solid', fgColor=hex6)
thin = Side(style='thin', color='FFD1D5DB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '代运营商达标分析'

COLS = ['代运营商','价格带','阈值','实付GTV','超价格带GTV','门店数','实际占比','达标状态','超出/剩余']
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
tc = ws.cell(1, 1, f'粤闽浙城市群 超价格带订单GTV占比 — 代运营商达标分析（{datetime.date.today().strftime("%Y-%m-%d")}）')
tc.fill = make_fill(CLR_HEADER)
tc.font = Font(bold=True, color='FFFFFFFF', size=12)
tc.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 28

for ci, col in enumerate(COLS, 1):
    c = ws.cell(2, ci, col)
    c.fill = make_fill(CLR_HEADER)
    c.font = Font(bold=True, color='FFFFFFFF', size=10)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = border
ws.row_dimensions[2].height = 22

for ri, row in enumerate(df_s1.itertuples(index=False), 3):
    band = row[9]
    is_total = row[10]
    colors = BAND_COLORS.get(band, BAND_COLORS['品类A'])

    bg = colors['total_bg'] if is_total else colors['row_bg']
    fg = colors['total_fg'] if is_total else colors['row_fg']

    for ci, val in enumerate(row[:9], 1):
        c = ws.cell(ri, ci, val)
        c.border = border
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.fill = make_fill(bg)
        c.font = Font(color=fg, bold=is_total, size=10)

        # 只对代运营商行（非合计行）做色阶
        if not is_total:
            status_val = str(row[7])
            if ci == 8:  # 达标状态
                if '超标' in status_val:
                    c.fill = make_fill(CLR_RED_BG)
                    c.font = Font(bold=True, color=CLR_RED_FG, size=10)
                else:
                    c.fill = make_fill(CLR_GRN_BG)
                    c.font = Font(bold=True, color=CLR_GRN_FG, size=10)
            elif ci == 7:  # 实际占比
                if '超标' in status_val:
                    c.fill = make_fill(CLR_RED_BG)
                    c.font = Font(bold=True, color=CLR_RED_FG, size=10)
            elif ci == 9:  # 超出/剩余
                try:
                    v = float(str(val).replace('%','').replace('+',''))
                    if v > 0:
                        c.fill = make_fill(CLR_RED_BG)
                        c.font = Font(bold=True, color=CLR_RED_FG, size=10)
                    elif v > -5:
                        c.fill = make_fill(CLR_YLW_BG)
                        c.font = Font(color=CLR_YLW_FG, size=10)
                except: pass

    ws.row_dimensions[ri].height = 20

col_widths = [14, 16, 7, 14, 14, 8, 10, 10, 10]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

import datetime; out = f'/tmp/超价格带分析_{datetime.date.today().strftime("%Y%m%d")}.xlsx'
wb.save(out)
print('saved:', out)
