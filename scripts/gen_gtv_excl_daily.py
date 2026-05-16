#!/usr/bin/env python3
"""
GTV剔除率分析每日自动发送脚本
每天9:30执行：生成GTV剔除率分析Excel并发给英杰(3264994429)
数据来源：/tmp/gtv_fraud_rate.xlsx（由fetch_finance_data.py生成）
"""
import subprocess, sys, os
from datetime import datetime
from pathlib import Path

DATE_STR = datetime.now().strftime("%Y-%m-%d")
INPUT_FILE = "/tmp/gtv_fraud_rate.xlsx"
OUTPUT_FILE = f"/tmp/GTV剔除率分析_{DATE_STR}.xlsx"
GEN_SCRIPT = "/tmp/gen_gtv_excl_excel.py"
TARGET_UID = "3264994429"

def log(msg):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)

# 1. 检查原始数据是否存在
if not Path(INPUT_FILE).exists():
    log(f"❌ 原始数据不存在: {INPUT_FILE}，尝试先抓取...")
    r = subprocess.run(["python3", "/root/.openclaw/workspace/scripts/fetch_finance_data.py"],
                       capture_output=True, text=True, timeout=120)
    log(r.stdout[-500:] if r.stdout else "no output")
    if not Path(INPUT_FILE).exists():
        log("❌ 抓取后仍无数据，退出")
        sys.exit(1)

log(f"✅ 原始数据存在: {INPUT_FILE}")

# 2. 生成Excel（inline，不依赖外部脚本文件）
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

def fill(color): return PatternFill("solid", fgColor=color)
def font(bold=False, size=10, color="000000", white=False):
    return Font(bold=bold, size=size, color="FFFFFF" if white else color)
def ctr(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
def left(): return Alignment(horizontal="left", vertical="center")
def bdr():
    s = Side(style="thin", color="DDDDDD")
    return Border(left=s, right=s, top=s, bottom=s)

F_HEADER = fill("1d4ed8")
F_RED    = fill("fee2e2")
F_YELLOW = fill("fef9c3")
F_GREEN  = fill("dcfce7")
F_TOTAL  = fill("FFF7E6")
F_WHITE  = fill("FFFFFF")

def excl_fill(rate):
    if rate >= 20: return F_RED
    if rate >= 15: return F_YELLOW
    return F_GREEN

def excl_font(rate):
    if rate >= 20: return font(color="b91c1c")
    if rate >= 15: return font(color="92400e")
    return font(color="166534")

ORDER = ['惠州市巨浪', '惠州市柏万', '泉州市小花猫', '惠安广通', '慈溪哲光', '义乌市汐洋']

wb_in = openpyxl.load_workbook(INPUT_FILE)
ws_in = wb_in.active
rows = []
for r in ws_in.iter_rows(min_row=2, values_only=True):
    if r[0] is None: continue
    org_raw = str(r[5]) if r[5] else ""
    org = next((k for k in ORDER if k in org_raw), None)
    if not org and len(r) > 12 and r[12]:
        org = next((k for k in ORDER if k in str(r[12])), None)
    if not org: continue
    gtv_raw = float(r[8] or 0)
    gtv_ver = float(r[10] or 0)
    excl_gtv = gtv_raw - gtv_ver
    excl_rate = round(excl_gtv / gtv_raw * 100, 2) if gtv_raw > 0 else 0
    rows.append({'org': org, 'shop': r[1], 'sale_id': r[2], 'sale_name': r[3],
                 'org_detail': r[5], 'cat1': r[6], 'cat2': r[7],
                 'gtv_raw': gtv_raw, 'gtv_ver': gtv_ver, 'excl_gtv': excl_gtv, 'excl_rate': excl_rate})

log(f"读取 {len(rows)} 条数据")

wb_out = openpyxl.Workbook()
ws1 = wb_out.active
ws1.title = "代运营商汇总"
ws1.merge_cells("A1:I1")
c = ws1["A1"]
c.value = f"粤闽浙城市群 实付验证GTV剔除率分析 — {DATE_STR}"
c.font = font(bold=True, size=12, white=True)
c.fill = F_HEADER
c.alignment = ctr()
ws1.row_dimensions[1].height = 28

headers = ["代运营商","门店总数","实付GTV(元)","实付验证GTV(元)","剔除GTV(元)","剔除率%","零GTV门店数","零GTV占比%","高剔除门店(>30%)"]
for ci, h in enumerate(headers, 1):
    c = ws1.cell(2, ci, h)
    c.font = font(bold=True, white=True); c.fill = F_HEADER; c.alignment = ctr(); c.border = bdr()
ws1.row_dimensions[2].height = 22

org_stats = {}
for op in ORDER:
    sub = [r for r in rows if r['org'] == op]
    shops = len(sub); gtv_raw = sum(r['gtv_raw'] for r in sub); gtv_ver = sum(r['gtv_ver'] for r in sub)
    excl_gtv = gtv_raw - gtv_ver; excl_rate = round(excl_gtv/gtv_raw*100, 2) if gtv_raw > 0 else 0
    zero_shops = sum(1 for r in sub if r['gtv_raw'] == 0)
    zero_pct = round(zero_shops/shops*100, 1) if shops > 0 else 0
    high_excl = sum(1 for r in sub if r['excl_rate'] > 30)
    org_stats[op] = {'shops':shops,'gtv_raw':gtv_raw,'gtv_ver':gtv_ver,'excl_gtv':excl_gtv,'excl_rate':excl_rate,'zero_shops':zero_shops,'zero_pct':zero_pct,'high_excl':high_excl}

for ri, op in enumerate(ORDER, 3):
    s = org_stats[op]
    for ci, v in enumerate([op,s['shops'],round(s['gtv_raw'],2),round(s['gtv_ver'],2),round(s['excl_gtv'],2),s['excl_rate'],s['zero_shops'],s['zero_pct'],s['high_excl']], 1):
        c = ws1.cell(ri, ci, v); c.alignment = ctr(); c.border = bdr()
        if ci == 6: c.fill = excl_fill(s['excl_rate']); c.font = excl_font(s['excl_rate'])
        else: c.fill = F_WHITE

tot_row = len(ORDER) + 3
tot_shops=sum(s['shops'] for s in org_stats.values()); tot_gtv_raw=sum(s['gtv_raw'] for s in org_stats.values())
tot_gtv_ver=sum(s['gtv_ver'] for s in org_stats.values()); tot_excl_gtv=tot_gtv_raw-tot_gtv_ver
tot_excl_rate=round(tot_excl_gtv/tot_gtv_raw*100,2) if tot_gtv_raw>0 else 0
tot_zero=sum(s['zero_shops'] for s in org_stats.values()); tot_zero_pct=round(tot_zero/tot_shops*100,1) if tot_shops>0 else 0
tot_high=sum(s['high_excl'] for s in org_stats.values())
for ci,v in enumerate(["合计",tot_shops,round(tot_gtv_raw,2),round(tot_gtv_ver,2),round(tot_excl_gtv,2),tot_excl_rate,tot_zero,tot_zero_pct,tot_high],1):
    c=ws1.cell(tot_row,ci,v); c.alignment=ctr(); c.border=bdr(); c.fill=F_TOTAL; c.font=font(bold=True)

note_row = tot_row+2
ws1.cell(note_row,1,"颜色说明：").font=font(bold=True)
for ci,(lbl,ff) in enumerate([("剔除率≥20% 超标",F_RED),("剔除率15-20% 临近",F_YELLOW),("剔除率<15% 达标",F_GREEN)],2):
    c=ws1.cell(note_row,ci,lbl); c.fill=ff; c.alignment=ctr()
for i,w in enumerate([18,10,16,18,16,10,12,12,14],1):
    ws1.column_dimensions[get_column_letter(i)].width=w

# Sheet2 品类分析
ws2 = wb_out.create_sheet("品类分析")
for ci,h in enumerate(["代运营商","一级品类","门店数","实付GTV(元)","实付验证GTV(元)","剔除率%","占该运营商GTV%"],1):
    c=ws2.cell(1,ci,h); c.font=font(bold=True,white=True); c.fill=F_HEADER; c.alignment=ctr(); c.border=bdr()
ri2=2
for op in ORDER:
    sub_op=[r for r in rows if r['org']==op]; op_gtv_ver=sum(r['gtv_ver'] for r in sub_op)
    cat_stats=defaultdict(lambda:{'shops':0,'gtv_raw':0,'gtv_ver':0})
    for r in sub_op:
        c1=r['cat1'] or '未知'; cat_stats[c1]['shops']+=1; cat_stats[c1]['gtv_raw']+=r['gtv_raw']; cat_stats[c1]['gtv_ver']+=r['gtv_ver']
    op_shown=False
    for cat,cs in sorted(cat_stats.items(),key=lambda x:-x[1]['gtv_ver']):
        excl=round((cs['gtv_raw']-cs['gtv_ver'])/cs['gtv_raw']*100,2) if cs['gtv_raw']>0 else 0
        pct_op=round(cs['gtv_ver']/op_gtv_ver*100,1) if op_gtv_ver>0 else 0
        for ci,v in enumerate([op if not op_shown else None,cat,cs['shops'],round(cs['gtv_raw'],2),round(cs['gtv_ver'],2),excl,pct_op],1):
            c=ws2.cell(ri2,ci,v); c.alignment=ctr(); c.border=bdr()
            if ci==6: c.fill=excl_fill(excl); c.font=excl_font(excl)
        op_shown=True; ri2+=1
for i,w in enumerate([18,14,10,16,18,10,14],1):
    ws2.column_dimensions[get_column_letter(i)].width=w

# Sheet3 高剔除门店
ws3=wb_out.create_sheet("高剔除门店(>30%)")
for ci,h in enumerate(["代运营商","门店名称","责任销售","组织名称","一级品类","二级品类","实付GTV","实付验证GTV","剔除率%"],1):
    c=ws3.cell(1,ci,h); c.font=font(bold=True,white=True); c.fill=F_HEADER; c.alignment=ctr(); c.border=bdr()
high_rows=sorted([r for r in rows if r['excl_rate']>30],key=lambda x:(-x['excl_rate'],x['org']))
for ri,r in enumerate(high_rows,2):
    for ci,v in enumerate([r['org'],r['shop'],r['sale_name'],r['org_detail'],r['cat1'],r['cat2'],round(r['gtv_raw'],2),round(r['gtv_ver'],2),r['excl_rate']],1):
        c=ws3.cell(ri,ci,v); c.alignment=left() if ci==2 else ctr(); c.border=bdr()
        if ci==9: c.fill=F_RED if r['excl_rate']>=50 else F_YELLOW; c.font=font(color="b91c1c")
for i,w in enumerate([14,28,10,30,10,12,12,14,10],1):
    ws3.column_dimensions[get_column_letter(i)].width=w

# Sheet4 原始明细
ws4=wb_out.create_sheet("原始明细")
for ci,h in enumerate(["验证商户ID","门店名称","责任销售工号","责任销售姓名","责任销售业绩归属组织ID","责任销售业绩归属组织名称","一级品类名称","二级品类名称","实付GTV金额","实付GTV金额——风控","实付验证GTV金额","实付验证GTV金额——风控+退款","代运营商","剔除率%"],1):
    c=ws4.cell(1,ci,h); c.font=font(bold=True,white=True); c.fill=F_HEADER; c.alignment=ctr(); c.border=bdr()
for ri,r in enumerate(rows,2):
    for ci,v in enumerate([None,r['shop'],r['sale_id'],r['sale_name'],None,r['org_detail'],r['cat1'],r['cat2'],round(r['gtv_raw'],2),0,round(r['gtv_ver'],2),0,r['org'],r['excl_rate']],1):
        c=ws4.cell(ri,ci,v); c.alignment=left() if ci==2 else ctr(); c.border=bdr()

wb_out.save(OUTPUT_FILE)
log(f"✅ Excel已生成: {OUTPUT_FILE} ({Path(OUTPUT_FILE).stat().st_size//1024}KB)")

# 3. 发送给英杰
log("📤 发送给英杰...")
r = subprocess.run(
    ["openclaw", "message", "send", "--target", TARGET_UID,
     "--channel", "daxiang", "--media", OUTPUT_FILE,
     "--message", f"GTV剔除率分析表（{DATE_STR}，昨日门店数据）\n📊 Sheet1: 代运营商汇总\n📊 Sheet2: 品类分析\n📊 Sheet3: 高剔除门店(>30%)\n📊 Sheet4: 原始明细"],
    capture_output=True, text=True, timeout=60
)
if r.returncode == 0:
    log("✅ 发送成功")
else:
    log(f"❌ 发送失败: {r.stderr[-300:]}")
    sys.exit(1)
