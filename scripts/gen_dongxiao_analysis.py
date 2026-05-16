import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import json, os, sys, re

INPUT = sys.argv[1] if len(sys.argv) > 1 else '/tmp/files/dongxiao_latest.xlsx'
OUTPUT = '/tmp/粤闽浙好店职人码动销分析表.xlsx'
HAODIAN_POOL = '/root/.openclaw/workspace/data/haodian_pool.xlsx'
ARCHIVE_DIR = '/root/.openclaw/workspace/data/zrq_orders_archive'
os.makedirs(ARCHIVE_DIR, exist_ok=True)

TARGETS = {'惠州巨浪':225,'惠州柏万':162,'泉州小花猫':193,'惠安广通':132,'慈溪哲光':151,'义乌汐洋':162}  # 5月播报组口径，2026-05-09英杰确认
ORDER = ['惠州巨浪','惠州柏万','泉州小花猫','惠安广通','慈溪哲光','义乌汐洋']

def get_org(s):
    s = str(s)
    if '巨浪' in s: return '惠州巨浪'
    if '柏万' in s: return '惠州柏万'
    if '小花猫' in s: return '泉州小花猫'
    if '广通' in s or '晋江' in s: return '惠安广通'
    if '哲光' in s: return '慈溪哲光'
    if '汐洋' in s: return '义乌汐洋'
    return None

def fill(hex): return PatternFill('solid', fgColor=hex)
def font(bold=False, color='000000', size=11): return Font(bold=bold, color=color, size=size)
def center(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def right(): return Alignment(horizontal='right', vertical='center')
def border():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)

# 读数据
df = pd.read_excel(INPUT)

# ── 自动识别格式 ──
NEW_FORMAT = '中小发代运营高线城市_优质动销门店数' in df.columns
# 第三种格式：职人码商户明细（门店码核销订单数字段，用好店底表匹配代运营商）
ZRQ_DETAIL_FORMAT = '门店码核销订单数(剔除风控+退款)' in df.columns

if ZRQ_DETAIL_FORMAT:
    # 用好店底表匹配代运营商
    df_pool = pd.read_excel(HAODIAN_POOL)
    pool_sale_map = {}
    for _, r in df_pool.iterrows():
        ag_name = None
        org = str(r.get('组织',''))
        for k,v in [('巨浪','惠州巨浪'),('柏万','惠州柏万'),('小花猫','泉州小花猫'),('广通','惠安广通'),('哲光','慈溪哲光'),('汐洋','义乌汐洋')]:
            if k in org: ag_name = v; break
        sale = str(r.get('跟进销售','')).split('/')[0].strip()
        if ag_name and sale and sale not in pool_sale_map:
            pool_sale_map[sale] = ag_name
    pool_shopid_map = {}
    for _, r in df_pool.iterrows():
        ag_name = None
        org = str(r.get('组织',''))
        for k,v in [('巨浪','惠州巨浪'),('柏万','惠州柏万'),('小花猫','泉州小花猫'),('广通','惠安广通'),('哲光','慈溪哲光'),('汐洋','义乌汐洋')]:
            if k in org: ag_name = v; break
        sid = str(r.get('shopId','')).strip()
        if ag_name and sid: pool_shopid_map[sid] = ag_name

    ID_COL = '门店ID'
    NAME_COL = '商户名称'
    SALES_COL = '销售工号名称'
    ORDER_COL = '门店码核销订单数(剔除风控+退款)'
    GTV_COL = '门店码实付验证GTV(剔除风控+退款)'
    df[ID_COL] = df[ID_COL].astype(str).str.strip()
    df['代运营商'] = df[ID_COL].map(pool_shopid_map)
    df['优惠码订单量'] = pd.to_numeric(df[ORDER_COL], errors='coerce').fillna(0)
    df['门店名称'] = df[NAME_COL]
    df['销售名称'] = df[SALES_COL]
    good_dedup = df[df['代运营商'].notna()].copy().drop_duplicates(ID_COL)
    DONE_THRESHOLD = 4
    m = re.search(r'(\d{4}-\d{2}-\d{2})', INPUT)
    data_date = m.group(1) if m else datetime.now().strftime('%Y-%m-%d')
    print(f"职人码商户明细格式，日期: {data_date}，好店匹配: {len(good_dedup)}家")
elif NEW_FORMAT:
    # 新格式：中小发代运营高线城市_优质动销门店数 >= 1 为达标
    DONG_COL = '中小发代运营高线城市_优质动销门店数'
    ID_COL = '门店id'
    NAME_COL = '门店名称'
    SALES_COL = '员工姓名'
    ORG_COL = '业绩归属组织名称'
    df['代运营商'] = df[ORG_COL].apply(get_org)
    df[ID_COL] = df[ID_COL].astype(str).str.strip()
    df[DONG_COL] = pd.to_numeric(df[DONG_COL], errors='coerce').fillna(0)
    # 新格式：全量门店即为好店池（已过滤好店），用动销数作为"订单量"等价字段
    df['优惠码订单量'] = df[DONG_COL]  # 0/1 标记，1=达标（等价≥4单）
    df['门店ID'] = df[ID_COL]
    df['销售名称'] = df[SALES_COL]
    good_dedup = df[df['代运营商'].notna()].copy().drop_duplicates('门店ID')
    # 新格式：达标阈值为1（不是4）
    DONE_THRESHOLD = 1
    # 日期：从文件名提取或用今天
    m = re.search(r'(\d{4}-\d{2}-\d{2})', INPUT)
    data_date = m.group(1) if m else datetime.now().strftime('%Y-%m-%d')
    print(f"新格式文件，日期: {data_date}，门店数: {len(good_dedup)}")
else:
    # 旧格式：是否代运营好店=1 + 优惠码订单量
    ID_COL = '门店ID'
    NAME_COL = '门店名称'
    SALES_COL = '销售名称'
    ORG_COL = '销售业绩归属组织名称'
    df['代运营商'] = df[ORG_COL].apply(get_org)
    df['门店ID'] = df['门店ID'].astype(str).str.strip()
    df['优惠码订单量'] = pd.to_numeric(df['优惠码订单量'], errors='coerce').fillna(0)
    good = df[df['是否代运营好店'] == 1].copy()
    good_dedup = good.sort_values('优惠码订单量', ascending=False).drop_duplicates('门店ID')
    data_date = df['日期分区字段，格式为date(yyyy-mm-dd)'].dropna().astype(str).max()[:10]
    print(f"旧格式文件，日期: {data_date}，好店数: {len(good_dedup)}")
    DONE_THRESHOLD = 4

# 读昨日存档（计算昨日净增）
yest_files = sorted([f for f in os.listdir(ARCHIVE_DIR) if f.endswith('.json') and f < f'{data_date}.json'])
yest_orders = {}
if yest_files:
    with open(os.path.join(ARCHIVE_DIR, yest_files[-1])) as f:
        yest_data = json.load(f)
        yest_orders = yest_data.get('orders', {})
    print(f"昨日存档: {yest_files[-1]}，记录数: {len(yest_orders)}")
else:
    print("无昨日存档，昨日净增显示0")

# 今日存档（统一用订单量/动销值存档）
today_orders = {row['门店ID']: int(row['优惠码订单量']) for _, row in good_dedup.iterrows()}
archive_path = os.path.join(ARCHIVE_DIR, f'{data_date}.json')
with open(archive_path, 'w') as f:
    json.dump({'date': data_date, 'orders': today_orders, 'format': 'new' if NEW_FORMAT else 'old'}, f, ensure_ascii=False)
print(f"已存档: {archive_path}")

wb = openpyxl.Workbook()

# ── Sheet1 代运营商纬度 ──
ws1 = wb.active
ws1.title = '代运营商纬度'

# 大标题
ws1.merge_cells('A1:M1')
c = ws1['A1']
c.value = f'粤闽浙好店职人码动销分析表（{data_date}）'
c.font = Font(bold=True, size=14, color='FFFFFF')
c.fill = fill('1E3A5F')
c.alignment = center()
ws1.row_dimensions[1].height = 30

headers = ['代运营商','好店池','5月目标','✅已达标','达标率','目标缺口',
           '📈日净增达标','🟡2-3单','🔴1单','❌零动销','已绑码','未绑码',
           '好店动销率']  # 13列，含日净增达标
for col, h in enumerate(headers, 1):
    c = ws1.cell(row=2, column=col, value=h)
    c.font = Font(bold=True, color='FFFFFF', size=10)
    c.fill = fill('2C5282')
    c.alignment = center()
    c.border = border()
ws1.row_dimensions[2].height = 40

for row_i, org in enumerate(ORDER, 3):
    sub = good_dedup[good_dedup['代运营商'] == org]
    
    pool = len(sub)
    tgt = TARGETS[org]
    ok = len(sub[sub['优惠码订单量'] >= DONE_THRESHOLD])
    rate = ok/tgt if tgt else 0
    gap = ok - tgt
    if NEW_FORMAT:
        s23 = 0; s1 = 0  # 新格式只有0/1，无中间单数
    else:
        s23 = len(sub[sub['优惠码订单量'].between(2,3)])
        s1 = len(sub[sub['优惠码订单量'] == 1])
    s0 = len(sub[sub['优惠码订单量'] == 0])
    bound = len(sub[sub['优惠码订单量'] > 0])
    unbound = pool - bound
    dong_rate = ok/pool if pool else 0
    
    # 昨日净增
    yest_ok_set = {k for k,v in yest_orders.items() if v >= DONE_THRESHOLD}
    today_ok_set = {row['门店ID'] for _, row in sub.iterrows() if row['优惠码订单量'] >= DONE_THRESHOLD}
    yest_dong_set = {k for k,v in yest_orders.items() if v > 0}
    today_dong_set = {row['门店ID'] for _, row in sub.iterrows() if row['优惠码订单量'] > 0}
    inc_ok = len(today_ok_set - yest_ok_set)
    inc_dong = len(today_dong_set - yest_dong_set)
    
    bg = 'EBF4FF' if row_i % 2 == 0 else 'FFFFFF'
    vals = [org, pool, tgt, ok, f'{rate*100:.1f}%', gap, inc_ok, s23, s1, s0, bound, unbound,
            f'{dong_rate*100:.1f}%']
    for col, v in enumerate(vals, 1):
        c = ws1.cell(row=row_i, column=col, value=v)
        c.fill = fill(bg)
        c.border = border()
        c.alignment = center() if col in (1,5,13) else right()
        # 达标率色阶
        if col == 5:
            pct = rate*100
            c.font = Font(bold=True, color='15803D' if pct>=80 else ('92400E' if pct>=60 else 'DC2626'))
        # 缺口
        if col == 6:
            c.font = Font(color='15803D' if gap>=0 else 'DC2626', bold=True)
        # 日净增达标色阶
        if col == 7:
            c.font = Font(bold=True, color='15803D' if inc_ok>0 else 'DC2626')

# 合计行
row = len(ORDER) + 3
totals = {}
for col_i, field in enumerate(['好店池','5月目标','✅已达标'], 2):
    totals[col_i] = sum(
        (len(good_dedup[good_dedup['代运营商']==o]) if field=='好店池' else
         TARGETS[o] if field=='5月目标' else
         len(good_dedup[(good_dedup['代运营商']==o)&(good_dedup['优惠码订单量']>=DONE_THRESHOLD)]))
        for o in ORDER
    )
total_rate = totals[4]/totals[3] if totals[3] else 0
total_gap = totals[4] - totals[3]

ws1.cell(row=row, column=1, value='粤闽浙合计').font = Font(bold=True, color='FFFFFF', size=11)
ws1.cell(row=row, column=1).fill = fill('1E3A5F')
ws1.cell(row=row, column=1).alignment = center()
ws1.cell(row=row, column=1).border = border()
for col_i in range(2, 14):
    c = ws1.cell(row=row, column=col_i)
    if col_i == 2: c.value = totals[2]
    elif col_i == 3: c.value = totals[3]
    elif col_i == 4: c.value = totals[4]
    elif col_i == 5: c.value = f'{total_rate*100:.1f}%'
    elif col_i == 6: c.value = total_gap
    # col 7 = 日净增合计（留空）
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = fill('1E3A5F')
    c.alignment = center()
    c.border = border()
ws1.row_dimensions[row].height = 24

# 列宽
widths = [14,8,8,8,8,8,10,8,8,8,8,8,10]
for i, w in enumerate(widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

# ── Sheet2 门店明细 ──
ws2 = wb.create_sheet('门店明细（全量好店）')
hdr2 = ['代运营商','门店ID','门店名称','销售名称','优惠码订单量','是否达标(≥4单)','昨日订单量','昨日净增']
for col, h in enumerate(hdr2, 1):
    c = ws2.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = fill('2C5282')
    c.alignment = center()
    c.border = border()

row = 2
for org in ORDER:
    sub = good_dedup[good_dedup['代运营商']==org].sort_values('优惠码订单量', ascending=False)
    for _, r in sub.iterrows():
        orders = int(r['优惠码订单量'])
        yest_ord = yest_orders.get(str(r['门店ID']), 0)
        ok_flag = '✅' if orders >= DONE_THRESHOLD else ''
        bg = 'F0FFF4' if orders >= DONE_THRESHOLD else ('FFFBEB' if (not NEW_FORMAT and orders in [2,3]) else ('FFF1F2' if (not NEW_FORMAT and orders == 1) else 'FFFFFF'))
        vals = [org, r['门店ID'], r['门店名称'], r['销售名称'], orders, ok_flag, yest_ord, orders-yest_ord]
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=row, column=col, value=v)
            c.fill = fill(bg[1:] if bg.startswith('#') else bg)
            c.border = border()
            c.alignment = Alignment(vertical='center')
        row += 1

# ── Sheet3 接近达标 ──
ws3 = wb.create_sheet('⭐接近达标（2-3单）')
for col, h in enumerate(hdr2, 1):
    c = ws3.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color='FFFFFF'); c.fill = fill('2C5282'); c.alignment = center(); c.border = border()
near = good_dedup[good_dedup['优惠码订单量'].between(2,3)].sort_values(['代运营商','优惠码订单量'], ascending=[True,False]) if not NEW_FORMAT else good_dedup[good_dedup['优惠码订单量']==0].head(0)  # 新格式无中间单数
for row_i, (_, r) in enumerate(near.iterrows(), 2):
    orders = int(r['优惠码订单量'])
    vals = [r['代运营商'], r['门店ID'], r['门店名称'], r['销售名称'], orders, '', yest_orders.get(str(r['门店ID']),0), '']
    for col, v in enumerate(vals, 1):
        c = ws3.cell(row=row_i, column=col, value=v)
        c.fill = fill('FFFBEB'); c.border = border(); c.alignment = Alignment(vertical='center')

# ── Sheet4 原始数据 ──
ws4 = wb.create_sheet('原始数据')
orig_cols = [c for c in df.columns]
for col, h in enumerate(orig_cols, 1):
    c = ws4.cell(row=1, column=col, value=h)
    c.font = Font(bold=True); c.fill = fill('EBF4FF'); c.border = border()
for row_i, (_, r) in enumerate(df.iterrows(), 2):
    for col, v in enumerate(r, 1):
        ws4.cell(row=row_i, column=col, value=v).border = border()

# ── Sheet5 日净增达标明细 ──
ws5 = wb.create_sheet('📈日净增达标明细')
hdr5 = ['代运营商','门店ID','门店名称','销售名称','今日订单量','昨日订单量','净增']
for col, h in enumerate(hdr5, 1):
    c = ws5.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = fill('15803D')
    c.alignment = center()
    c.border = border()
    ws5.column_dimensions[get_column_letter(col)].width = [14,12,22,12,12,12,8][col-1]
ws5.row_dimensions[1].height = 24

row5 = 2
for org in ORDER:
    sub = good_dedup[good_dedup['代运营商']==org]
    for _, r in sub.iterrows():
        sid = str(r['门店ID'])
        today_ord = int(r['优惠码订单量'])
        yest_ord  = yest_orders.get(sid, 0)
        inc = today_ord - yest_ord
        # 日净增达标：今日≥4，昨日<4
        if today_ord >= DONE_THRESHOLD and yest_ord < DONE_THRESHOLD:
            vals = [org, r['门店ID'], r['门店名称'], r['销售名称'], today_ord, yest_ord, inc]
            for col, v in enumerate(vals, 1):
                c = ws5.cell(row=row5, column=col, value=v)
                c.fill = fill('F0FFF4')
                c.border = border()
                c.alignment = center() if col in (1,5,6,7) else Alignment(vertical='center')
            row5 += 1

if row5 == 2:
    ws5.cell(row=2, column=1, value='（暂无日净增达标门店）').font = Font(italic=True, color='888888')

# ── Sheet6 销售维度 ──
ws6 = wb.create_sheet('销售维度')
hdr6 = ['代运营商','销售名称','好店池','✅已达标','达标率','📈日净增达标','目标缺口','🟡2-3单','🔴1单','❌零动销']
for col, h in enumerate(hdr6, 1):
    c = ws6.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = fill('2C5282')
    c.alignment = center()
    c.border = border()
    ws6.column_dimensions[get_column_letter(col)].width = [14,12,8,10,8,12,10,10,8,10][col-1]
ws6.row_dimensions[1].height = 28

AGENT_COLORS6 = {
    '惠州巨浪':'FFF2CC','惠州柏万':'DDEBF7','泉州小花猫':'E2EFDA',
    '惠安广通':'FCE4D6','慈溪哲光':'EAE1F4','义乌汐洋':'FDEBD0'
}
row6 = 2
for org in ORDER:
    sub_org = good_dedup[good_dedup['代运营商']==org]
    sellers = sub_org['销售名称'].unique()
    # 代运营商分组标题
    ws6.merge_cells(f'A{row6}:J{row6}')
    gc = ws6.cell(row=row6, column=1, value=f'▶ {org}')
    gc.font = Font(bold=True, color='FFFFFF', size=10)
    gc.fill = fill('E07B39')
    gc.alignment = Alignment(horizontal='left', vertical='center')
    gc.border = border()
    ws6.row_dimensions[row6].height = 20
    row6 += 1

    seller_rows = []
    for seller in sellers:
        sub_s = sub_org[sub_org['销售名称']==seller]
        pool_s = len(sub_s)
        ok_s   = len(sub_s[sub_s['优惠码订单量'] >= DONE_THRESHOLD])
        # 日净增
        inc_s = sum(
            1 for _, r in sub_s.iterrows()
            if int(r['优惠码订单量']) >= DONE_THRESHOLD and yest_orders.get(str(r['门店ID']),0) < DONE_THRESHOLD
        )
        gap_s  = ok_s - int(pool_s * TARGETS.get(org,0) / max(len(sub_org),1))
        s23_s  = len(sub_s[sub_s['优惠码订单量'].between(2,3)]) if not NEW_FORMAT else 0
        s1_s   = len(sub_s[sub_s['优惠码订单量']==1]) if not NEW_FORMAT else 0
        s0_s   = len(sub_s[sub_s['优惠码订单量']==0])
        rate_s = ok_s/pool_s if pool_s else 0
        seller_rows.append((seller, pool_s, ok_s, rate_s, inc_s, gap_s, s23_s, s1_s, s0_s))

    # 按已达标降序
    seller_rows.sort(key=lambda x: x[2], reverse=True)
    bg_color = AGENT_COLORS6.get(org,'FFFFFF')
    for i, (seller, pool_s, ok_s, rate_s, inc_s, gap_s, s23_s, s1_s, s0_s) in enumerate(seller_rows):
        bg = bg_color if i%2==0 else 'FFFFFF'
        vals6 = [org, seller, pool_s, ok_s, f'{rate_s*100:.1f}%', inc_s, gap_s, s23_s, s1_s, s0_s]
        for col, v in enumerate(vals6, 1):
            c = ws6.cell(row=row6, column=col, value=v)
            c.fill = fill(bg)
            c.border = border()
            c.alignment = center() if col in (1,2,5) else Alignment(horizontal='right', vertical='center')
            if col == 5:
                pct = rate_s*100
                c.font = Font(bold=True, color='15803D' if pct>=80 else ('92400E' if pct>=60 else 'DC2626'))
            if col == 6:
                c.font = Font(bold=True, color='15803D' if inc_s>0 else '888888')
        row6 += 1

ws6.freeze_panes = 'A2'

wb.save(OUTPUT)
print(f'Done: {OUTPUT}')
print(f'\n代运营商达标汇总:')
for org in ORDER:
    sub = good_dedup[good_dedup['代运营商']==org]
    ok = len(sub[sub['优惠码订单量']>=DONE_THRESHOLD])
    tgt = TARGETS[org]
    print(f'  {org}: {ok}/{tgt} ({ok/tgt*100:.0f}%)')
